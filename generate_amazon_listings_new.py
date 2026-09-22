"""
Amazon Bulk Listing Generator
==============================
Generates one Amazon-upload-ready .xlsm file per phone model, using:
  1. BASE_TEMPLATE      - a real Amazon category workbook (all sheets/validation kept intact)
  2. models_master.xlsx - one row per model, every field that is the SAME across all
                          designs of that model (bullets, price, keywords, dimensions, ...)
  3. images_master.xlsx - one row per (model, design number), the main image URL for that design

Columns are located by their Amazon field code in row 5 of the Template sheet
(e.g. "item_name[marketplace_id=...]#1.value"), NOT by hardcoded position. Amazon
reorders columns between template versions, so position-based mapping silently
writes data into the wrong fields.

Run:
    python generate_amazon_listings2.py

Output:
    outputs/<model_key>.xlsm   (one file per model row in models_master.xlsx)
"""

import io
import os
import re
import sys
import warnings

import openpyxl
import pandas as pd

warnings.simplefilter("ignore")

# ---------- CONFIG ----------
BASE_TEMPLATE_PATH = "iPhone_17-processing-summary-processing-summary.xlsm"
MODELS_MASTER_PATH = "models_master_new.xlsx"
IMAGES_MASTER_PATH = "images_master.xlsx"
OUTPUT_DIR = "outputs"
TEMPLATE_SHEET_NAME = "Template"
FIELD_CODE_ROW = 5          # row 5 holds Amazon's machine field codes
DATA_START_ROW = 7          # rows 1-6 are Amazon headers/examples
CHUNK_SIZE_DEFAULT = 24     # designs per parent SKU

# Logical name -> Amazon field code, normalized by stripping every [bracket] group so the
# same map works across marketplaces and template versions.
FIELDS = {
    "sku":                   "contribution_sku#1.value",
    "product_type":          "product_type#1.value",
    "listing_action":        "::record_action",
    "parentage_level":       "parentage_level#1.value",
    "parent_sku":            "child_parent_sku_relationship#1.parent_sku",
    "variation_theme":       "variation_theme#1.name",
    "item_name":             "item_name#1.value",
    "brand_name":            "brand#1.value",
    "product_id_type":       "amzn1.volt.ca.product_id_type",
    "browse_node":           "recommended_browse_nodes#1.value",
    "manufacturer":          "manufacturer#1.value",
    "main_image_url":        "main_product_image_locator#1.media_location",
    "other_image_url_1":     "other_product_image_locator_1#1.media_location",
    "other_image_url_2":     "other_product_image_locator_2#1.media_location",
    "other_image_url_3":     "other_product_image_locator_3#1.media_location",
    "product_description":   "product_description#1.value",
    "bullet_point_1":        "bullet_point#1.value",
    "bullet_point_2":        "bullet_point#2.value",
    "bullet_point_3":        "bullet_point#3.value",
    "bullet_point_4":        "bullet_point#4.value",
    "generic_keyword":       "generic_keyword#1.value",
    "number_of_items":       "number_of_items#1.value",
    "item_type_name":        "item_type_name#1.value",
    "color":                 "color#1.value",
    "manufacturer_contact":  "rtip_manufacturer_contact_information#1.value",
    "compatible_devices":    "compatible_devices#1.value",
    "unit_count":            "unit_count#1.value",
    "unit_count_type":       "unit_count#1.type.value",
    "ext_prod_info_entity":  "external_product_information#1.entity",
    "ext_prod_info_value":   "external_product_information#1.value",
    "packer_contact":        "packer_contact_information#1.value",
    "item_length":           "item_length_width#1.length.value",
    "item_length_unit":      "item_length_width#1.length.unit",
    "item_width":            "item_length_width#1.width.value",
    "item_width_unit":       "item_length_width#1.width.unit",
    "item_weight":           "item_weight#1.value",
    "item_weight_unit":      "item_weight#1.unit",
    "condition_type":        "condition_type#1.value",
    "tax_code":              "product_tax_code#1.value",
    "fulfillment_channel":   "fulfillment_availability#1.fulfillment_channel_code",
    "quantity":              "fulfillment_availability#1.quantity",
    "price":                 "purchasable_offer#1.our_price#1.schedule#1.value_with_tax",
    "mrp":                   "purchasable_offer#1.maximum_retail_price#1.schedule#1.value_with_tax",
    "pkg_length":            "item_package_dimensions#1.length.value",
    "pkg_length_unit":       "item_package_dimensions#1.length.unit",
    "pkg_width":             "item_package_dimensions#1.width.value",
    "pkg_width_unit":        "item_package_dimensions#1.width.unit",
    "pkg_height":            "item_package_dimensions#1.height.value",
    "pkg_height_unit":       "item_package_dimensions#1.height.unit",
    "pkg_weight":            "item_package_weight#1.value",
    "pkg_weight_unit":       "item_package_weight#1.unit",
    "country_of_origin":     "country_of_origin#1.value",
    "warranty_description":  "warranty_description#1.value",
    "batteries_required":    "batteries_required#1.value",
    "dangerous_goods":       "supplier_declared_dg_hz_regulation#1.value",
}

# Amazon writes these back when it processes an upload. Any value inherited from the
# base template is a stale result for a previous feed, so they are always cleared.
STATUS_FIELD_CODES = [
    "::submission_status",
    "::number_of_attributes_with_errors",
    "::number_of_attributes_with_other_suggestions",
]

# Columns read from models_master.xlsx. Everything else on a row is derived:
# sku, item_name, color, parentage_level, parent_sku, main_image_url.
MODEL_FIELDS = [
    "product_type", "listing_action", "variation_theme", "brand_name", "product_id_type",
    "browse_node", "manufacturer", "other_image_url_1", "other_image_url_2",
    "other_image_url_3", "product_description", "bullet_point_1", "bullet_point_2",
    "bullet_point_3", "bullet_point_4", "generic_keyword", "number_of_items",
    "item_type_name", "manufacturer_contact", "compatible_devices", "unit_count",
    "unit_count_type", "ext_prod_info_entity", "ext_prod_info_value", "packer_contact",
    "item_length", "item_length_unit", "item_width", "item_width_unit", "item_weight",
    "item_weight_unit", "condition_type", "tax_code", "fulfillment_channel", "quantity",
    "price", "mrp", "pkg_length", "pkg_length_unit", "pkg_width", "pkg_width_unit",
    "pkg_height", "pkg_height_unit", "pkg_weight", "pkg_weight_unit",
    "country_of_origin", "warranty_description", "batteries_required", "dangerous_goods",
]

# Exact parent/child field split taken from the known-good iPhone_17 upload.
# A parent row is a variation hub: it carries the shared marketing copy and nothing
# offer-specific (no price, quantity, images, color or parent_sku).
PARENT_MODEL_FIELDS = [
    "product_type", "listing_action", "variation_theme", "brand_name",
    "product_description", "bullet_point_1", "bullet_point_2", "bullet_point_3",
    "bullet_point_4", "generic_keyword", "manufacturer_contact", "packer_contact",
    "country_of_origin", "batteries_required", "dangerous_goods",
]
# batteries_required appears on parents only in the known-good file.
CHILD_EXCLUDED_FIELDS = {"batteries_required"}

REQUIRED_MODEL_COLS = ["model_key", "sku_prefix", "title_prefix"]
NUM_DESIGNS_FALLBACK_COL = "num_designs"


def model_specific_defaults(model_key):
    """Fields that must never inherit another model's text if left blank."""
    name = model_key.replace("_", " ").strip()
    return {
        "compatible_devices": name,
        "generic_keyword": f"{name} skin, {name} wrap, {name} sticker, mobile skin for {name}",
        "bullet_point_4": f"{name} skin, {name} wrap, {name} sticker, matte skin for {name}",
    }


def normalize_code(code):
    """'item_name[marketplace_id=X][language_tag=en_IN]#1.value' -> 'item_name#1.value'"""
    return re.sub(r"\[[^\]]*\]", "", str(code)).strip()


def build_column_map(ws):
    """normalized field code -> 0-indexed column, read from row 5 of the Template sheet."""
    codes = [c.value for c in ws[FIELD_CODE_ROW]]
    by_code = {}
    for i, raw in enumerate(codes):
        if raw in (None, ""):
            continue
        by_code.setdefault(normalize_code(raw), i)  # first occurrence wins

    col, missing = {}, []
    for name, code in FIELDS.items():
        if code in by_code:
            col[name] = by_code[code]
        else:
            missing.append(name)

    status_cols = [by_code[c] for c in STATUS_FIELD_CODES if c in by_code]
    return col, missing, status_cols


def read_models_master(path):
    df = pd.read_excel(path, dtype=str)
    df.columns = [str(c).strip() for c in df.columns]
    # A duplicated header in the sheet makes pandas emit 'model_key.1'; drop those.
    df = df.loc[:, ~df.columns.str.match(r".*\.\d+$")]
    for req in REQUIRED_MODEL_COLS:
        if req not in df.columns:
            raise ValueError(f"{path} is missing required column: {req}")
    df = df[df["model_key"].notna()].reset_index(drop=True)
    df["model_key"] = df["model_key"].astype(str).str.strip()
    return df


def read_images_master(path):
    df = pd.read_excel(path, dtype=str)
    df.columns = [str(c).strip().lower() for c in df.columns]
    required = {"model_key", "design_number", "main_image_url"}
    if not required.issubset(set(df.columns)):
        raise ValueError(f"{path} must have columns: {sorted(required)}")
    df = df[df["model_key"].notna() & df["design_number"].notna()].copy()
    df["model_key"] = df["model_key"].astype(str).str.strip()
    df["design_number"] = df["design_number"].astype(float).astype(int)
    return df.sort_values("design_number")


def check_model_key_alignment(models_df, images_df):
    """Warn (with typo suggestions) when a model has no images and no num_designs fallback."""
    if NUM_DESIGNS_FALLBACK_COL in models_df.columns:
        has_fallback = models_df[NUM_DESIGNS_FALLBACK_COL].notna() & (
            models_df[NUM_DESIGNS_FALLBACK_COL].astype(str).str.strip() != ""
        )
    else:
        has_fallback = pd.Series(False, index=models_df.index)

    image_keys = set(images_df["model_key"])
    missing = set(models_df.loc[~has_fallback, "model_key"]) - image_keys
    if not missing:
        return

    print("\n" + "=" * 70)
    print("WARNING: model_key mismatch between models_master and images_master")
    print("=" * 70)
    for mk in sorted(missing):
        print(f"    '{mk}' has no image rows")
        close = [ik for ik in image_keys if ik.lower() == mk.lower()]
        if close:
            print(f"        -> did you mean '{close[0]}' (different case)?")
    print(f"\n  model_key values present in images_master: {sorted(image_keys)}")
    print("=" * 70 + "\n")


def cell_value(model_row, field, defaults):
    """models_master value if filled, else a derived model-specific default, else None.

    Deliberately does NOT fall back to the base template's own data rows: those belong
    to whichever model the template was last used for, so inheriting them silently
    stamps that model's copy onto a different product.
    """
    raw = model_row.get(field)
    if raw is not None and not pd.isna(raw) and str(raw).strip() != "":
        return str(raw).strip()
    return defaults.get(field)


def build_title(title_prefix, design_number):
    return f"{str(title_prefix).rstrip()} D-{design_number}"


def resolve_design_list(model_row, images_df):
    """(design_numbers, image_urls, images_missing) for one model."""
    model_key = model_row["model_key"]
    rows = images_df[images_df["model_key"] == model_key]
    if not rows.empty:
        return rows["design_number"].tolist(), rows["main_image_url"].tolist(), False

    raw = model_row.get(NUM_DESIGNS_FALLBACK_COL)
    if raw is None or pd.isna(raw) or str(raw).strip() == "":
        raise ValueError(
            f"no rows in {IMAGES_MASTER_PATH} for model_key='{model_key}', and no "
            f"'{NUM_DESIGNS_FALLBACK_COL}' in {MODELS_MASTER_PATH} either. Run "
            f"upload_images.py for this model, or set {NUM_DESIGNS_FALLBACK_COL} to "
            f"generate with blank image URLs."
        )
    n = int(float(raw))
    return list(range(1, n + 1)), [None] * n, True


def generate_model_workbook(base_bytes, model_row, images_df, chunk_size, col, status_cols):
    wb = openpyxl.load_workbook(io.BytesIO(base_bytes), keep_vba=True)
    ws = wb[TEMPLATE_SHEET_NAME]

    model_key = model_row["model_key"]
    sku_prefix = str(model_row["sku_prefix"]).strip()
    title_prefix = model_row["title_prefix"]
    defaults = model_specific_defaults(model_key)

    design_numbers, image_urls, images_missing = resolve_design_list(model_row, images_df)
    n_designs = len(design_numbers)
    n_parents = (n_designs + chunk_size - 1) // chunk_size

    values = {f: cell_value(model_row, f, defaults) for f in MODEL_FIELDS}
    blanks = sorted(f for f, v in values.items() if v is None)

    # Clear the old data region by value rather than delete_rows(): deleting rows on this
    # sheet discards the workbook's data validation ranges and embedded media.
    n_cols = ws.max_column
    for r in range(DATA_START_ROW, ws.max_row + 1):
        for c in range(1, n_cols + 1):
            ws.cell(row=r, column=c).value = None

    def write(row, field, value):
        if field in col and value is not None:
            ws.cell(row=row, column=col[field] + 1, value=value)

    current_row = DATA_START_ROW
    parent_skus = [f"{sku_prefix}_parent_{k + 1}" for k in range(n_parents)]

    for k, parent_sku in enumerate(parent_skus):
        write(current_row, "sku", parent_sku)
        write(current_row, "parentage_level", "Parent")
        write(current_row, "item_name", build_title(title_prefix, design_numbers[k * chunk_size]))
        for field in PARENT_MODEL_FIELDS:
            write(current_row, field, values[field])
        current_row += 1

    for idx, design_n in enumerate(design_numbers):
        write(current_row, "sku", f"{sku_prefix}_D{design_n}")
        write(current_row, "parentage_level", "Child")
        write(current_row, "parent_sku", parent_skus[idx // chunk_size])
        write(current_row, "item_name", build_title(title_prefix, design_n))
        write(current_row, "color", f"Design-{design_n}")
        write(current_row, "main_image_url", image_urls[idx])
        for field in MODEL_FIELDS:
            if field not in CHILD_EXCLUDED_FIELDS:
                write(current_row, field, values[field])
        current_row += 1

    # Amazon's own result columns, if this template came from a processing report.
    for r in range(DATA_START_ROW, current_row):
        for c in status_cols:
            ws.cell(row=r, column=c + 1).value = None

    return wb, n_designs, n_parents, images_missing, blanks


def verify_workbook(path, n_designs, n_parents, sku_prefix, col):
    """Re-open the saved file and assert the structure Amazon accepted."""
    wb = openpyxl.load_workbook(path)
    ws = wb[TEMPLATE_SHEET_NAME]

    def get(r, f):
        return ws.cell(row=r, column=col[f] + 1).value

    end = DATA_START_ROW + n_parents + n_designs
    levels = [get(r, "parentage_level") for r in range(DATA_START_ROW, end)]
    assert levels[:n_parents] == ["Parent"] * n_parents, "parent rows missing"
    assert levels[n_parents:] == ["Child"] * n_designs, "child rows missing"

    first_child = DATA_START_ROW + n_parents
    assert get(DATA_START_ROW, "sku") == f"{sku_prefix}_parent_1"
    assert get(DATA_START_ROW, "parent_sku") is None, "parent row must not have a parent_sku"
    assert get(DATA_START_ROW, "price") is None, "parent row must not carry a price"
    assert get(first_child, "parent_sku") == f"{sku_prefix}_parent_1"
    child_design = str(get(first_child, "sku")).split("_D")[-1]
    assert get(first_child, "color") == f"Design-{child_design}", "color/SKU design mismatch"
    for f in ("price", "mrp", "quantity", "item_name", "product_type"):
        assert get(first_child, f) not in (None, ""), f"child missing {f}"
    wb.close()


def resolve_base_template_path():
    if os.path.exists(BASE_TEMPLATE_PATH):
        return BASE_TEMPLATE_PATH
    candidates = [f for f in os.listdir(".") if f.lower().endswith(".xlsm")]
    preferred = [f for f in candidates if "processing-summary" in f.lower()]
    if len(preferred) == 1:
        print(f"NOTE: '{BASE_TEMPLATE_PATH}' not found; using '{preferred[0]}'")
        return preferred[0]
    if len(candidates) == 1:
        print(f"NOTE: '{BASE_TEMPLATE_PATH}' not found; using '{candidates[0]}'")
        return candidates[0]
    if not candidates:
        sys.exit(
            "ERROR: base template not found. Put your Amazon template .xlsm next to "
            "this script, or set BASE_TEMPLATE_PATH."
        )
    sys.exit(
        f"ERROR: '{BASE_TEMPLATE_PATH}' not found and multiple .xlsm files exist "
        f"({', '.join(candidates)}). Set BASE_TEMPLATE_PATH to the right one."
    )


def main():
    base_template_path = resolve_base_template_path()
    for p in (MODELS_MASTER_PATH, IMAGES_MASTER_PATH):
        if not os.path.exists(p):
            sys.exit(f"ERROR: not found: {p}")

    os.makedirs(OUTPUT_DIR, exist_ok=True)

    with open(base_template_path, "rb") as f:
        base_bytes = f.read()

    probe = openpyxl.load_workbook(io.BytesIO(base_bytes), keep_vba=True)
    if TEMPLATE_SHEET_NAME not in probe.sheetnames:
        sys.exit(f"ERROR: '{base_template_path}' has no '{TEMPLATE_SHEET_NAME}' sheet.")
    col, missing, status_cols = build_column_map(probe[TEMPLATE_SHEET_NAME])
    probe.close()

    print(f"Base template : {base_template_path}")
    print(f"Mapped        : {len(col)}/{len(FIELDS)} fields by Amazon field code")
    if missing:
        print(f"  not in this template (skipped): {', '.join(missing)}")
    for req in ("sku", "parentage_level", "item_name", "main_image_url"):
        if req not in col:
            sys.exit(f"ERROR: required field '{req}' not found in {base_template_path}.")

    models_df = read_models_master(MODELS_MASTER_PATH)
    images_df = read_images_master(IMAGES_MASTER_PATH)
    print(f"Models        : {len(models_df)} from {MODELS_MASTER_PATH}")
    check_model_key_alignment(models_df, images_df)

    for _, model_row in models_df.iterrows():
        model_key = model_row["model_key"]
        raw_chunk = model_row.get("chunk_size")
        chunk_size = (
            int(float(raw_chunk))
            if raw_chunk is not None and not pd.isna(raw_chunk) and str(raw_chunk).strip() != ""
            else CHUNK_SIZE_DEFAULT
        )

        try:
            wb, n_designs, n_parents, images_missing, blanks = generate_model_workbook(
                base_bytes, model_row, images_df, chunk_size, col, status_cols
            )
        except ValueError as e:
            print(f"  [SKIPPED] {model_key}: {e}")
            continue

        out_path = os.path.join(OUTPUT_DIR, f"{model_key}.xlsm")
        wb.save(out_path)
        wb.close()
        verify_workbook(out_path, n_designs, n_parents, str(model_row["sku_prefix"]).strip(), col)

        print(f"  [OK] {model_key}: {n_designs} designs, {n_parents} parent SKU(s) -> {out_path}")
        if images_missing:
            print("        NOTE: main_image_url left blank on every row - fill before uploading")
        if blanks:
            print(f"        blank in models_master ({len(blanks)}): {', '.join(blanks)}")

    print("Done.")


if __name__ == "__main__":
    main()
